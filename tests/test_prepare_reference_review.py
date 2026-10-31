"""Tests of the reference review script (Issue #46), on fictitious tickets."""

from __future__ import annotations

import importlib.util
import random
import subprocess
import sys
from pathlib import Path

import pytest

SCRIPT = Path(__file__).resolve().parent.parent / "scripts" / "prepare_reference_review.py"
spec = importlib.util.spec_from_file_location("prepare_reference_review", SCRIPT)
review = importlib.util.module_from_spec(spec)
spec.loader.exec_module(review)

QUOTAS = {"accounts": 6, "network": 5}


def fake_row(ident, category, rule, body):
    return {
        "id": ident,
        "subject": "",
        "body": f"orig {body}",
        "body_fr": body,
        "original_labels": {"label": "x"},
        "proposed_labels": {
            "category": category,
            "rule": rule,
            "type": None,
            "doubtful": ident.endswith("0"),
        },
    }


def fake_rows():
    rows = []
    for category in ("accounts", "network"):
        for rule in ("r1", "r2"):
            for i in range(10):
                body = f"{category} {rule} ticket {'abcdefghij'[i] * 12} unique texte numéro {i}"
                rows.append(fake_row(f"{category}-{rule}-{i}", category, rule, body))
    rows.append(fake_row("FORCED", "network", "r1", "ticket forcé tout à fait particulier"))
    return rows


def selected_ids(seed):
    return [r["id"] for r in review.select(fake_rows(), random.Random(seed), QUOTAS, "FORCED")]


def test_selection_respects_quotas_and_forces_the_given_ticket():
    picked = review.select(fake_rows(), random.Random(1), QUOTAS, "FORCED")
    counts = {c: sum(r["proposed_labels"]["category"] == c for r in picked) for c in QUOTAS}
    assert counts == QUOTAS
    assert "FORCED" in {r["id"] for r in picked}
    assert len({r["id"] for r in picked}) == len(picked)


def test_selection_rotates_over_rules():
    picked = review.select(fake_rows(), random.Random(1), QUOTAS, None)
    rules = {r["proposed_labels"]["rule"] for r in picked if r["id"].startswith("accounts")}
    assert rules == {"r1", "r2"}


def test_selection_skips_near_duplicates_and_invents_nothing():
    rows = [
        fake_row(
            f"d{i}",
            "accounts",
            "r1",
            "même début de ticket très long, assez pour dépasser la clé " + "x" * i,
        )
        for i in range(6)
    ]
    rows.append(fake_row("other", "accounts", "r1", "tout autre sujet de ticket différent"))
    picked = review.select(rows, random.Random(1), {"accounts": 5}, None)
    assert len([r for r in picked if r["id"].startswith("d")]) == 1  # one per family
    assert len(picked) == 2  # quota not reachable without duplicates


def test_selection_is_reproducible_and_depends_on_the_seed():
    assert selected_ids(7) == selected_ids(7)
    assert selected_ids(7) != selected_ids(8)


def test_default_quotas_make_fifty_tickets():
    assert sum(review.QUOTAS.values()) == 50


def test_script_refuses_to_overwrite_an_existing_file(tmp_path):
    existing = tmp_path / "relecture.xlsx"
    existing.write_bytes(b"review in progress")
    result = subprocess.run(
        [sys.executable, str(SCRIPT), "--output", str(existing), "--source", str(tmp_path / "x")],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert result.returncode == 1
    assert "Refus" in result.stdout
    assert existing.read_bytes() == b"review in progress"


def test_workbook_has_the_columns_to_fill_and_validations(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    picked = review.select(fake_rows(), random.Random(1), QUOTAS, "FORCED")
    taxonomy = {
        "categories": {"accounts": "Comptes", "network": "Réseau"},
        "types": {"incident": "Incident"},
        "priorities": {"P1": "Critique"},
    }
    path = tmp_path / "out.xlsx"
    review.build_workbook(picked, taxonomy).save(path)
    book = load_workbook(path)
    sheet = book["Relecture"]
    assert sheet.max_row == len(picked) + 1
    assert [sheet.cell(row=1, column=c).value for c in (9, 10)] == [
        "Type (à remplir)",
        "Priorité (à remplir)",
    ]
    assert sheet.cell(row=2, column=9).value is None  # nothing pre-filled for the developer
    assert len(sheet.data_validations.dataValidation) == 3
    assert "Aide" in book.sheetnames
