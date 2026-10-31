"""Build the reference review workbook: 50 translated test tickets for the developer to check.

The developer reads each ticket, confirms or corrects the category and fills the type and the
priority, which the external test set almost never provides. Those 50 tickets are the only real
tickets whose type and priority are verified (Issue #46).

Usage:
    uv run --group data python scripts/prepare_reference_review.py
    uv run --group data python scripts/prepare_reference_review.py --output other.xlsx

The script REFUSES to overwrite an existing file (the developer fills it in by hand): pass
``--force`` only to rebuild a file that holds no review yet.
"""

from __future__ import annotations

import argparse
import json
import random
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

import yaml

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.data.corpus import template_key  # noqa: E402
from src.data.label_mapping import load_mapping  # noqa: E402

SEED = 42
QUOTAS = {
    "accounts": 8,
    "messaging": 7,
    "software": 7,
    "printing": 6,
    "infrastructure": 5,
    "telephony": 5,
    "network": 4,
    "security": 4,
    "workstation": 4,
}
FORCED_ID = "PUB-ITS-00323"  # the only ticket whose tag was not restored by the translation
DEFAULT_SOURCE = ROOT / "data" / "raw" / "public" / "it_support_fr.jsonl"
DEFAULT_TAXONOMY = ROOT / "config" / "taxonomy.yaml"
DEFAULT_OUTPUT = ROOT / "data" / "raw" / "reference" / "relecture_50_tickets.xlsx"


def select(
    rows: list[dict[str, Any]],
    rng: random.Random,
    quotas: dict[str, int] = QUOTAS,
    forced_id: str | None = FORCED_ID,
) -> list[dict[str, Any]]:
    """Pick tickets per category, rotating over rules and skipping near-duplicates.

    Each row needs ``id``, ``body_fr`` and ``proposed_labels`` (category and rule). The forced
    ticket, if it exists and its category has a quota, is always part of the selection.
    """
    by_category: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        by_category[row["proposed_labels"]["category"]].append(row)

    forced = next((row for row in rows if row["id"] == forced_id), None)
    picked: list[dict[str, Any]] = []
    for category, quota in quotas.items():
        chosen = [forced] if forced and forced["proposed_labels"]["category"] == category else []
        seen = {template_key(row["body_fr"]) for row in chosen}
        by_rule: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in by_category[category]:
            if row["id"] != forced_id:
                by_rule[row["proposed_labels"]["rule"]].append(row)
        for candidates in by_rule.values():
            rng.shuffle(candidates)
        # Rotate over rules until the quota is filled; near-duplicates are skipped.
        rules = sorted(by_rule)
        while len(chosen) < quota and any(by_rule[rule] for rule in rules):
            for rule in rules:
                while by_rule[rule] and len(chosen) < quota:
                    candidate = by_rule[rule].pop()
                    key = template_key(candidate["body_fr"])
                    if key not in seen:
                        seen.add(key)
                        chosen.append(candidate)
                        break
        picked.extend(chosen)
    rng.shuffle(picked)
    return picked


def with_current_labels(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Recompute ``proposed_labels`` from the current mapping rules (no need to translate again)."""
    mapping = load_mapping()
    return [
        {
            **row,
            "proposed_labels": mapping.propose(
                "it_support", row["original_labels"], row["subject"], row["body"]
            ).as_dict(),
        }
        for row in rows
    ]


def build_workbook(picked: list[dict[str, Any]], taxonomy: dict[str, Any]):
    """Excel workbook with one sheet to fill in and one help sheet (needs ``openpyxl``)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    sheet = book.active
    sheet.title = "Relecture"
    headers = [
        ("N°", 5),
        ("Identifiant", 15),
        ("Ticket traduit en français", 70),
        ("Texte d'origine", 45),
        ("Catégorie proposée", 15),
        ("Douteuse ?", 9),
        ("Type proposé", 15),
        ("Catégorie correcte (seulement si fausse)", 17),
        ("Type (à remplir)", 15),
        ("Priorité (à remplir)", 9),
        ("Traduction incompréhensible ? (x)", 13),
        ("Commentaire", 30),
    ]
    to_fill = PatternFill("solid", fgColor="FFF2CC")
    for column, (title, width) in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[cell.column_letter].width = width
    for index, row in enumerate(picked, start=1):
        labels = row["proposed_labels"]
        values = [
            index,
            row["id"],
            row["body_fr"],
            row["body"],
            labels["category"],
            "oui" if labels["doubtful"] else "",
            labels["type"] or "",
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=index + 1, column=column, value=value).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
        for column in range(8, 13):
            sheet.cell(row=index + 1, column=column).fill = to_fill

    last = len(picked) + 1
    for column, values in (
        ("H", taxonomy["categories"]),
        ("I", taxonomy["types"]),
        ("J", taxonomy["priorities"]),
    ):
        validation = DataValidation(
            type="list", formula1=f'"{",".join(values)}"', allow_blank=True, showErrorMessage=True
        )
        sheet.add_data_validation(validation)
        validation.add(f"{column}2:{column}{last}")
    sheet.freeze_panes = "C2"

    help_sheet = book.create_sheet("Aide")
    lines = [
        ("Comment relire", ""),
        ("1", "Lire le ticket. Le texte d'origine sert seulement si la traduction est obscure."),
        ("2", "Catégorie : ne rien écrire si la proposition est juste, sinon choisir la bonne."),
        ("3", "Type et priorité : à remplir pour chaque ticket (ces étiquettes manquent au jeu)."),
        ("4", "Sans information suffisante : choisir au mieux et le dire en commentaire."),
        ("5", "Compter environ 30 secondes par ticket, soit moins de 30 minutes en tout."),
        ("", ""),
        ("Catégories", ""),
        *taxonomy["categories"].items(),
        ("", ""),
        ("Types", ""),
        *taxonomy["types"].items(),
        ("", ""),
        ("Priorités", ""),
        *taxonomy["priorities"].items(),
    ]
    for number, (left, right) in enumerate(lines, start=1):
        help_sheet.cell(row=number, column=1, value=left)
        help_sheet.cell(row=number, column=2, value=str(right))
        if not right and left:
            help_sheet.cell(row=number, column=1).font = Font(bold=True)
    help_sheet.column_dimensions["A"].width = 18
    help_sheet.column_dimensions["B"].width = 90
    return book


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE, help="it_support_fr.jsonl")
    parser.add_argument("--taxonomy", type=Path, default=DEFAULT_TAXONOMY)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--force", action="store_true", help="overwrite an existing output file")
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")

    if args.output.exists() and not args.force:
        print(
            f"Refus d'écraser {args.output} : relecture peut-être en cours (--force pour forcer)."
        )
        return 1

    lines = args.source.read_text(encoding="utf-8").splitlines()
    rows = with_current_labels([json.loads(line) for line in lines if line.strip()])
    picked = select(rows, random.Random(SEED))
    taxonomy = yaml.safe_load(args.taxonomy.read_text(encoding="utf-8"))
    book = build_workbook(picked, taxonomy)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    book.save(args.output)

    counts: dict[str, int] = defaultdict(int)
    for row in picked:
        counts[row["proposed_labels"]["category"]] += 1
    doubtful = sum(row["proposed_labels"]["doubtful"] for row in picked)
    print(f"{len(picked)} tickets : {dict(counts)}, douteux : {doubtful}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
